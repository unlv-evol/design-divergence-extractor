"""Pilot workbooks and agreement, through dde-prep."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from dde_prep.cli import main

PACKET = """# Case {cid}

## File `src/A.java`

### Class `A`

### Class `A.Inner`
"""


def test_sheets_and_agreement(tmp_path: Path) -> None:
    packets = tmp_path / "packets"
    for cid in ("C1", "C2", "C3"):
        (packets / cid).mkdir(parents=True)
        (packets / cid / "packet.md").write_text(PACKET.format(cid=cid))
    books = {}
    for coder in ("coderA", "coderB"):
        books[coder] = tmp_path / f"{coder}.xlsx"
        assert (
            main(
                [
                    "pilot-sheets",
                    "pilot",
                    "--coder",
                    coder,
                    "--packets",
                    str(packets),
                    "--out",
                    str(books[coder]),
                ]
            )
            == 0
        )
    wb = load_workbook(books["coderA"])
    assert [wb["Pass1_Cases"][f"A{r}"].value for r in (2, 3, 4)] == ["C1", "C2", "C3"]
    assert wb["Pass1_Roles"]["C3"].value == "A.Inner"
    answers = {"coderA": ["design", "non_design", "both"], "coderB": ["design", "non_design", "none"]}
    for coder, cats in answers.items():
        wb = load_workbook(books[coder])
        ws = wb["Pass1_Cases"]
        for r, cat in enumerate(cats, start=2):
            ws[f"L{r}"] = cat
        ws["B2"], ws["K2"] = "Y", "yes"
        wb.save(books[coder])
    out = tmp_path / "agreement"
    assert (
        main(["pilot-agreement", "--a", str(books["coderA"]), "--b", str(books["coderB"]), "--out", str(out)])
        == 0
    )
    report = (out / "agreement_report.md").read_text()
    assert "Cases coded by both: 3" in report
    with (out / "codes_long.csv").open() as fh:
        codes = list(csv.DictReader(fh))
    assert {"case_id": "C1", "coder": "coderA", "code": "responsibility_split_merge"} in codes
