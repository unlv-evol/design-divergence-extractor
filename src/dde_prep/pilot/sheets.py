"""Build coding workbooks for the pilot.

  practice:  dde-prep pilot-sheets practice --coder coderA --packets ../practice/packets --out practice_coderA.xlsx
  pilot:     dde-prep pilot-sheets pilot --coder coderA --packets pilot_packets/ \
                 [--cases pilot_cases.csv] --out coding_coderA.xlsx

Case rows and one role row per edited class are prefilled from the packets (``## File`` and
``### Class`` headings). Without packets, the sheets are left empty for the coder to fill.
Needs openpyxl.
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

DESIGN = [
    "responsibility_split_merge",
    "cross_layer_relocation",
    "missing_replaced_abstraction",
    "api_contract_redesign",
    "collaboration_changed",
]
CONTROL = ["refactoring_present", "feature_absence", "independent_edit_same_lines", "dependency_mismatch"]
PASS2_ONLY = [
    "refactoring_unsupported_by_repatch",
    "refactoring_missed_by_rminer",
    "conflict_non_mo_file",
    "tool_failure",
    "unexplained",
]
ROLES = [
    "Information Holder",
    "Structurer",
    "Service Provider",
    "Controller",
    "Coordinator",
    "Interfacer",
    "unclear",
]
CASE_COLS = (
    ["case_id"]
    + DESIGN
    + CONTROL
    + [
        "design_affects_patch",
        "overall_category",
        "suggested_category",
        "confidence",
        "evidence",
        "new_code_candidate",
        "memo",
        "minutes",
    ]
)
ROLE_COLS = ["case_id", "file", "class", "source_role", "target_role", "role_changed", "note"]
PASS2_COLS = [
    "case_id",
    "conflict_stage",
    "file",
    "conflict_ref",
    "attributed_code",
    "links_to_pass1",
    "note",
]
MAX_ROWS = 200

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
DESIGN_FILL = PatternFill("solid", fgColor="2E75B6")
CONTROL_FILL = PatternFill("solid", fgColor="7F6000")
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")
CALC_FILL = PatternFill("solid", fgColor="EDEDED")
WARN_FILL = PatternFill("solid", fgColor="F8CBAD")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")


def col(i: int) -> str:
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def parse_packets(
    packets: Path | None, cases: list[str] | None
) -> tuple[list[str], list[tuple[str, str, str]]]:
    if packets is None:
        return cases or [], []
    ids = cases or sorted(p.parent.name for p in packets.glob("*/packet.md"))
    roles = []
    for cid in ids:
        text = (packets / cid / "packet.md").read_text(encoding="utf-8")
        current = ""
        for line in text.splitlines():
            m = re.match(r"## File `(.+)`", line)
            if m:
                current = m.group(1)
            m = re.match(r"### Class `(.+)`", line)
            if m:
                roles.append((cid, current, m.group(1)))
    return ids, roles


def header(ws: Worksheet, cols: list[str], widths: dict[str, int]) -> None:
    for i, name in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = BOLD
        c.fill = DESIGN_FILL if name in DESIGN else CONTROL_FILL if name in CONTROL else HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[col(i)].width = widths.get(name, 14)
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "B2"


def dv(ws: Worksheet, values: list[str] | str, ref: str, prompt: str = "") -> None:
    formula = values if isinstance(values, str) else '"' + ",".join(values) + '"'
    v = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Not in the list",
        error="Pick a value from the dropdown.",
    )
    if prompt:
        v.promptTitle, v.prompt, v.showInputMessage = "How to fill", prompt, True
    ws.add_data_validation(v)
    v.add(ref)


def style_rows(ws: Worksheet, ncols: int, calc_cols: set[int]) -> None:
    for r in range(2, MAX_ROWS + 2):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = FONT
            cell.alignment = WRAP
            cell.border = THIN
            cell.fill = CALC_FILL if c in calc_cols else INPUT_FILL


def case_sheet(wb: Workbook, title: str, case_ids: list[str]) -> None:
    ws = wb.create_sheet(title)
    header(
        ws,
        CASE_COLS,
        {
            "case_id": 22,
            "evidence": 45,
            "new_code_candidate": 30,
            "memo": 40,
            "overall_category": 15,
            "suggested_category": 15,
            "design_affects_patch": 15,
            "minutes": 9,
            "confidence": 11,
            **{c: 12 for c in DESIGN + CONTROL},
        },
    )
    idx = {n: i + 1 for i, n in enumerate(CASE_COLS)}
    d0, d1 = col(idx[DESIGN[0]]), col(idx[DESIGN[-1]])
    c0, c1 = col(idx[CONTROL[0]]), col(idx[CONTROL[-1]])
    k, s = col(idx["design_affects_patch"]), col(idx["suggested_category"])
    oc = col(idx["overall_category"])
    style_rows(ws, len(CASE_COLS), {idx["suggested_category"]})
    for r in range(2, MAX_ROWS + 2):
        if r - 2 < len(case_ids):
            ws.cell(row=r, column=1, value=case_ids[r - 2])
        ws[f"{s}{r}"] = (
            f'=IF(A{r}="","",IF(AND(COUNTIF({d0}{r}:{d1}{r},"Y")>0,{k}{r}="yes"),'
            f'IF(COUNTIF({c0}{r}:{c1}{r},"Y")>0,"both","design"),'
            f'IF(COUNTIF({c0}{r}:{c1}{r},"Y")>0,"non_design","none")))'
        )
    last = MAX_ROWS + 1
    dv(ws, ["Y"], f"{d0}2:{c1}{last}", "Y if the code applies; leave blank if not.")
    dv(ws, ["yes", "nearby_only", "n/a"], f"{k}2:{k}{last}", "Codebook 6.2")
    dv(ws, ["design", "non_design", "both", "none"], f"{oc}2:{oc}{last}", "Codebook 6.1")
    conf = col(idx["confidence"])
    dv(ws, ["high", "medium", "low"], f"{conf}2:{conf}{last}")
    ws.conditional_formatting.add(
        f"{oc}2:{oc}{last}", FormulaRule(formula=[f'AND({oc}2<>"",{oc}2<>{s}2)'], fill=WARN_FILL)
    )


def role_sheet(wb: Workbook, title: str, rows: list[tuple[str, str, str]]) -> None:
    ws = wb.create_sheet(title)
    header(
        ws,
        ROLE_COLS,
        {
            "case_id": 22,
            "file": 50,
            "class": 24,
            "source_role": 18,
            "target_role": 18,
            "role_changed": 14,
            "note": 45,
        },
    )
    style_rows(ws, len(ROLE_COLS), {6})
    for r in range(2, MAX_ROWS + 2):
        if r - 2 < len(rows):
            for c, v in enumerate(rows[r - 2], start=1):
                ws.cell(row=r, column=c, value=v)
        ws[f"F{r}"] = (
            f'=IF(OR(D{r}="",E{r}=""),"",IF(OR(D{r}="unclear",E{r}="unclear"),"unclear",'
            f'IF(E{r}="no_counterpart","no_counterpart",IF(D{r}=E{r},"no","yes"))))'
        )
    last = MAX_ROWS + 1
    dv(ws, "=Lists!$A$2:$A$8", f"D2:D{last}", "Codebook section 7")
    dv(ws, "=Lists!$B$2:$B$9", f"E2:E{last}", "Codebook section 7; no_counterpart if none located")


def pass2_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Pass2_Conflicts")
    header(
        ws,
        PASS2_COLS,
        {
            "case_id": 22,
            "file": 50,
            "conflict_ref": 16,
            "attributed_code": 30,
            "conflict_stage": 15,
            "links_to_pass1": 14,
            "note": 45,
        },
    )
    ws["A1"].comment = None
    style_rows(ws, len(PASS2_COLS), set())
    last = MAX_ROWS + 1
    dv(ws, ["cherry_pick", "after_repatch"], f"B2:B{last}")
    dv(
        ws,
        f"=Lists!$C$2:$C${1 + len(DESIGN + CONTROL + PASS2_ONLY)}",
        f"E2:E{last}",
        "One code per conflict: design, control, or a pass-2 code (codebook 5.2)",
    )
    dv(ws, ["yes", "partly", "no"], f"F2:F{last}", "Did your pass-1 codes predict this conflict?")


def lists_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lists")
    ws["A1"], ws["B1"], ws["C1"] = "source_role", "target_role", "pass2_code"
    for i, r in enumerate(ROLES, start=2):
        ws[f"A{i}"] = r
        ws[f"B{i}"] = r
    ws[f"B{len(ROLES) + 2}"] = "no_counterpart"
    for i, c in enumerate(DESIGN + CONTROL + PASS2_ONLY, start=2):
        ws[f"C{i}"] = c
    ws.sheet_state = "hidden"


def legend(wb: Workbook, coder: str, kind: str) -> None:
    ws = wb.active
    ws.title = "Legend"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    lines = [
        (
            f"Pilot coding workbook — {coder} — {'practice cases' if kind == 'practice' else 'pilot cases'}",
            "",
        ),
        ("", ""),
        (
            "Fill in",
            "Yellow cells only. Grey cells are computed. Use the dropdowns; do not rename sheets or columns.",
        ),
        (
            "Pass1_Cases",
            "One row per case. Y for each code that applies (blue headers = design codes, brown = control codes).",
        ),
        (
            "Pass1_Roles",
            "One row per edited class. Pick the role stereotype on each side; role_changed is computed.",
        ),
        (
            "Pass2_Conflicts",
            "Only after the supervisor releases conflict output (Mon 2 Nov). One row per conflicting file."
            if kind == "pilot"
            else "Not used for practice.",
        ),
        (
            "suggested_category",
            "Computed from your codes (codebook 6.1). If your overall_category differs it turns orange; explain in memo.",
        ),
        ("", ""),
        ("Example row (worked example W1)", ""),
        ("case_id", "W1"),
        (
            "codes marked Y",
            "responsibility_split_merge, api_contract_redesign, refactoring_present, dependency_mismatch",
        ),
        ("design_affects_patch", "yes"),
        ("overall_category", "both"),
        ("confidence", "high"),
        (
            "evidence",
            "retryOrder moved to RetryPolicy (Move Method); find returns Optional<Order>; flush renamed flushAll; audit() absent in target",
        ),
        ("minutes", "25"),
        ("Roles example", "W1 · OrderService · Service Provider → Service Provider (role_changed = no)"),
        ("", ""),
        ("Codes", "Definitions are in the codebook (docs/02_codebook_v0.pdf); this list is a reminder only."),
    ]
    short = {
        "responsibility_split_merge": "Code the patch fixes is held by another class, split, or merged in the target",
        "cross_layer_relocation": "Edit site sits in another module or layer in the target",
        "missing_replaced_abstraction": "A mediating superclass/interface/helper is missing or replaced",
        "api_contract_redesign": "Edited method/class makes a different promise (signature, return, null, exceptions)",
        "collaboration_changed": "Edited class works with different collaborators, connected to the edit site",
        "refactoring_present": "A behaviour-preserving restructuring explains a difference at the edit site",
        "feature_absence": "The target never had the code the patch fixes, and nothing plays its role",
        "independent_edit_same_lines": "The target changed the same lines for its own reasons",
        "dependency_mismatch": "The patch uses something the target lacks, not by design",
    }
    lines += [(k, v) for k, v in short.items()]
    for i, (a, b) in enumerate(lines, start=1):
        ws[f"A{i}"], ws[f"B{i}"] = a, b
        ws[f"A{i}"].font = Font(
            name="Arial", size=12 if i == 1 else 10, bold=bool(a) and (i == 1 or b == "" or i <= 7)
        )
        ws[f"B{i}"].font = FONT
        ws[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A3"].fill = INPUT_FILL


def build(kind: str, coder: str, packets: Path | None, cases_csv: Path | None, out: Path) -> None:
    cases = None
    if cases_csv:
        with cases_csv.open(newline="", encoding="utf-8") as fh:
            cases = [r["case_id"].strip() for r in csv.DictReader(fh) if r["case_id"].strip()]
    ids, roles = parse_packets(packets, cases)
    if kind == "practice":
        ids = [i for i in ids if not i.startswith("W")]
        roles = [r for r in roles if not r[0].startswith("W")]
    wb = Workbook()
    legend(wb, coder, kind)
    case_sheet(wb, "Pass1_Cases", ids)
    role_sheet(wb, "Pass1_Roles", roles)
    if kind == "pilot":
        pass2_sheet(wb)
    lists_sheet(wb)
    wb.properties.creator = "Beyond Refactorings pilot"
    wb.properties.title = f"{kind} coding — {coder}"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}: {len(ids)} case rows, {len(roles)} role rows")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("kind", choices=["practice", "pilot"])
    ap.add_argument("--coder", required=True)
    ap.add_argument("--packets", type=Path)
    ap.add_argument("--cases", type=Path, help="pilot_cases.csv; default: every packet folder")
    ap.add_argument("--out", type=Path, required=True)
    a = ap.parse_args(argv)
    build(a.kind, a.coder, a.packets, a.cases, a.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
