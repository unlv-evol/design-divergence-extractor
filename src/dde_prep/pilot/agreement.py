"""Agreement between Coder A and Coder B, and export of codes for ``dde validate``.

  dde-prep pilot-agreement --a coderA/coding_coderA.xlsx --b coderB/coding_coderB.xlsx --out agreement/
  dde-prep pilot-agreement ... --consensus coding_consensus.xlsx     # also export consensus codes

Writes to --out:
  agreement_report.md   Checkpoint 2 result, Cohen's kappa per field and per code, roles, pass 2
  disagreements.csv     every disagreement, with an empty resolution column for the meeting
  codes_long.csv        case_id, coder, code: input to `dde validate --codes` (coder = consensus)

Reads only the input columns, never the cached formula values, so it works on files saved by
any spreadsheet program. Needs openpyxl.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

DESIGN = [
    "responsibility_split_merge",
    "cross_layer_relocation",
    "missing_replaced_abstraction",
    "api_contract_redesign",
    "collaboration_changed",
]
CONTROL = ["refactoring_present", "feature_absence", "independent_edit_same_lines", "dependency_mismatch"]
CODES = DESIGN + CONTROL
KAPPA_THRESHOLD = 0.6


def rows(path: Path, sheet: str) -> list[dict[str, str]]:
    wb = load_workbook(path, read_only=True)
    if sheet not in wb.sheetnames:
        return []
    it = wb[sheet].iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for r in it:
        d = {h: ("" if v is None else str(v).strip()) for h, v in zip(head, r, strict=False) if h}
        if d.get("case_id"):
            out.append(d)
    return out


def suggested(r: dict[str, str]) -> str:
    design = any(r.get(c) == "Y" for c in DESIGN) and r.get("design_affects_patch") == "yes"
    control = any(r.get(c) == "Y" for c in CONTROL)
    return "both" if design and control else "design" if design else "non_design" if control else "none"


def role_changed(r: dict[str, str]) -> str:
    s, t = r.get("source_role", ""), r.get("target_role", "")
    if not s or not t:
        return ""
    if "unclear" in (s, t):
        return "unclear"
    if t == "no_counterpart":
        return "no_counterpart"
    return "no" if s == t else "yes"


def kappa(a: list[str], b: list[str]) -> tuple[float | None, float, int]:
    """Cohen's kappa, observed agreement, n. Kappa is None when chance agreement is 1."""
    n = len(a)
    if n == 0:
        return None, 0.0, 0
    po = sum(x == y for x, y in zip(a, b, strict=False)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return (None if pe == 1 else (po - pe) / (1 - pe)), po, n


def landis_koch(k: float | None) -> str:
    if k is None:
        return "n/a"
    for bound, label in (
        (0, "poor"),
        (0.2, "slight"),
        (0.4, "fair"),
        (0.6, "moderate"),
        (0.8, "substantial"),
        (1.01, "almost perfect"),
    ):
        if k <= bound:
            return label
    return "almost perfect"


def fmt(k: float | None) -> str:
    return "n/a" if k is None else f"{k:.2f}"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--a", type=Path, required=True)
    ap.add_argument("--b", type=Path, required=True)
    ap.add_argument("--consensus", type=Path)
    ap.add_argument("--out", type=Path, default=Path("agreement"))
    args = ap.parse_args(argv)
    args.out.mkdir(parents=True, exist_ok=True)

    A = {r["case_id"]: r for r in rows(args.a, "Pass1_Cases")}
    B = {r["case_id"]: r for r in rows(args.b, "Pass1_Cases")}
    common = sorted(set(A) & set(B))
    only = sorted(set(A) ^ set(B))
    incomplete = [c for c in common if not A[c].get("overall_category") or not B[c].get("overall_category")]
    done = [c for c in common if c not in incomplete]

    report = [
        "# Pilot agreement, pass 1\n",
        f"Coder A: `{args.a.name}` (sha256 {sha(args.a)}); Coder B: `{args.b.name}` (sha256 {sha(args.b)}).",
        f"Cases coded by both: {len(done)}; incomplete: {len(incomplete)}; coded by only one: {len(only)}.\n",
    ]
    disagreements: list[dict[str, str]] = []

    # Checkpoint 2: overall category
    ka, po, n = kappa([A[c]["overall_category"] for c in done], [B[c]["overall_category"] for c in done])
    collapse = {"none": "non_design"}
    kc, poc, _ = kappa(
        [collapse.get(A[c]["overall_category"], A[c]["overall_category"]) for c in done],
        [collapse.get(B[c]["overall_category"], B[c]["overall_category"]) for c in done],
    )
    verdict = "PASS" if ka is not None and ka >= KAPPA_THRESHOLD else "FAIL"
    report += [
        "## Checkpoint 2: design vs non-design vs both\n",
        f"**{verdict}** at kappa >= {KAPPA_THRESHOLD} (proposed criterion; the team confirms).\n",
        "| Measure | Kappa | Band | Agreement | n |",
        "| --- | --- | --- | --- | --- |",
        f"| overall_category as coded (4 values) | {fmt(ka)} | {landis_koch(ka)} | {po:.0%} | {n} |",
        f"| with `none` merged into `non_design` (3 values) | {fmt(kc)} | {landis_koch(kc)} | {poc:.0%} | {n} |",
        "",
        "Confusion (rows Coder A, columns Coder B):\n",
    ]
    cats = ["design", "non_design", "both", "none"]
    conf = Counter((A[c]["overall_category"], B[c]["overall_category"]) for c in done)
    report += ["| A \\ B | " + " | ".join(cats) + " |", "|---" * (len(cats) + 1) + "|"]
    report += [f"| {x} | " + " | ".join(str(conf[(x, y)]) for y in cats) + " |" for x in cats]
    mismatch = [c for c in done for r in (A[c], B[c]) if r["overall_category"] != suggested(r)]
    if mismatch:
        report.append(
            f"\nCases where a coder's category differs from the rule in codebook 6.1: "
            f"{', '.join(sorted(set(mismatch)))} (check the memo)."
        )

    # Per code and per field
    report += [
        "\n## Codes\n",
        "| Code | Kappa | Band | Both Y | A only | B only | Neither |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for code in CODES:
        a = [A[c].get(code, "") == "Y" for c in done]
        b = [B[c].get(code, "") == "Y" for c in done]
        k, _, _ = kappa([str(x) for x in a], [str(y) for y in b])
        cnt = Counter(zip(a, b, strict=False))
        report.append(
            f"| `{code}` | {fmt(k)} | {landis_koch(k)} | {cnt[(True, True)]} | "
            f"{cnt[(True, False)]} | {cnt[(False, True)]} | {cnt[(False, False)]} |"
        )
    report += [
        "",
        "Kappa is n/a when both coders gave the same answer for every case; read the counts.",
        "\n## Other pass-1 fields\n",
        "| Field | Kappa | Agreement |",
        "| --- | --- | --- |",
    ]
    for field in ("design_affects_patch", "confidence"):
        k, p, _ = kappa([A[c].get(field, "") for c in done], [B[c].get(field, "") for c in done])
        report.append(f"| `{field}` | {fmt(k)} | {p:.0%} |")

    for c in done:
        for field in ["overall_category", "design_affects_patch"] + CODES:
            va, vb = A[c].get(field, ""), B[c].get(field, "")
            if va != vb:
                disagreements.append(
                    {
                        "case_id": c,
                        "item": field,
                        "coder_a": va or "(blank)",
                        "coder_b": vb or "(blank)",
                        "resolution": "",
                        "note": "",
                    }
                )

    # Roles
    RA = {(r["case_id"], r.get("file", ""), r.get("class", "")): r for r in rows(args.a, "Pass1_Roles")}
    RB = {(r["case_id"], r.get("file", ""), r.get("class", "")): r for r in rows(args.b, "Pass1_Roles")}
    keys = sorted(k for k in set(RA) & set(RB) if k[0] in done)
    report += [
        "\n## Role stereotypes\n",
        "| Field | Kappa | Band | Agreement | n |",
        "| --- | --- | --- | --- | --- |",
    ]
    for field, get in (
        ("source_role", lambda r: r.get("source_role", "")),
        ("target_role", lambda r: r.get("target_role", "")),
        ("role_changed", role_changed),
    ):
        k, p, n2 = kappa([get(RA[x]) for x in keys], [get(RB[x]) for x in keys])
        report.append(f"| `{field}` | {fmt(k)} | {landis_koch(k)} | {p:.0%} | {n2} |")
        for x in keys:
            if get(RA[x]) != get(RB[x]):
                disagreements.append(
                    {
                        "case_id": x[0],
                        "item": f"{field}:{x[2]}",
                        "coder_a": get(RA[x]) or "(blank)",
                        "coder_b": get(RB[x]) or "(blank)",
                        "resolution": "",
                        "note": "",
                    }
                )

    # Pass 2
    PA = {
        (r["case_id"], r.get("conflict_stage", ""), r.get("file", ""), r.get("conflict_ref", "")): r
        for r in rows(args.a, "Pass2_Conflicts")
    }
    PB = {
        (r["case_id"], r.get("conflict_stage", ""), r.get("file", ""), r.get("conflict_ref", "")): r
        for r in rows(args.b, "Pass2_Conflicts")
    }
    if PA or PB:
        k2 = sorted(set(PA) & set(PB))
        k, p, n2 = kappa(
            [PA[x].get("attributed_code", "") for x in k2], [PB[x].get("attributed_code", "") for x in k2]
        )
        report += [
            "\n## Pass 2: conflict attribution\n",
            f"Conflicts attributed by both: {n2} (only one coder: {len(set(PA) ^ set(PB))}). "
            f"Kappa {fmt(k)} ({landis_koch(k)}), agreement {p:.0%}.",
        ]
        dist = Counter(PA[x].get("attributed_code", "") for x in k2) + Counter(
            PB[x].get("attributed_code", "") for x in k2
        )
        report += ["", "| Attributed code | Times chosen (both coders) |", "| --- | --- |"]
        report += [f"| `{c}` | {n3} |" for c, n3 in dist.most_common()]

    if only or incomplete:
        report += ["\n## To complete\n"]
        report += [f"- coded by only one coder: {c}" for c in only]
        report += [f"- overall_category missing: {c}" for c in incomplete]

    # codes_long.csv
    long_rows = []
    sources = [("coderA", args.a), ("coderB", args.b)] + (
        [("consensus", args.consensus)] if args.consensus else []
    )
    for coder, path in sources:
        roles_changed = {r["case_id"] for r in rows(path, "Pass1_Roles") if role_changed(r) == "yes"}
        for r in rows(path, "Pass1_Cases"):
            codes = [c for c in CODES if r.get(c) == "Y"] + (
                ["role_changed"] if r["case_id"] in roles_changed else []
            )
            for c in codes or [""]:
                long_rows.append({"case_id": r["case_id"], "coder": coder, "code": c})

    (args.out / "agreement_report.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    for name, data, cols in (
        ("disagreements.csv", disagreements, ["case_id", "item", "coder_a", "coder_b", "resolution", "note"]),
        ("codes_long.csv", long_rows, ["case_id", "coder", "code"]),
    ):
        with (args.out / name).open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
            w.writeheader()
            w.writerows(data)
    print(
        f"Checkpoint 2: {verdict} (kappa {fmt(ka)}, n={n}); {len(disagreements)} disagreements; "
        f"report in {args.out / 'agreement_report.md'}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
